import os
import json
from pathlib import Path
from PIL import Image

try:
    import docx
    from docx.shared import Inches, Pt
except ImportError:
    docx = None

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, Table, TableStyle
except ImportError:
    SimpleDocTemplate = None

class CourtReportEngine:
    def __init__(self, analysis_dir="analysis", original_image_path=None):
        self.analysis_dir = Path(analysis_dir)
        self.original_image_path = original_image_path
        self.crops_dir = self.analysis_dir / "crops"
        self.reports_dir = self.analysis_dir / "reports"
        
        self.crops_dir.mkdir(parents=True, exist_ok=True)
        self.reports_dir.mkdir(parents=True, exist_ok=True)

    def generate_reports(self):
        dossier_path = self.analysis_dir / "forensic_dossier.json"
        if not dossier_path.exists():
            return None
        
        with open(dossier_path, "r") as f:
            dossier = json.load(f)
            
        # 1. Filter out un-localized findings
        valid_findings = []
        for item in dossier:
            bbox = item.get("bbox", {})
            w = bbox.get("width", 0)
            h = bbox.get("height", 0)
            if w > 0 and h > 0:
                valid_findings.append(item)
                
        # 2. Crop Screenshots
        self._generate_crops(valid_findings)
        
        # 3. Generate JSON Reports
        self._generate_json(valid_findings)
        
        # 4. Generate PDF Reports
        if SimpleDocTemplate:
            self._generate_pdf(valid_findings, "Executive")
            self._generate_pdf(valid_findings, "Technical")
            self._generate_pdf(valid_findings, "Court")
            
        # 5. Generate DOCX Reports
        if docx:
            self._generate_docx(valid_findings, "Executive")
            self._generate_docx(valid_findings, "Technical")
            self._generate_docx(valid_findings, "Court")
            
        return valid_findings

    def _generate_crops(self, findings):
        if not self.original_image_path or not os.path.exists(self.original_image_path):
            return
            
        try:
            img = Image.open(self.original_image_path)
            for item in findings:
                bbox = item.get("bbox", {})
                x = bbox.get("x", 0)
                y = bbox.get("y", 0)
                w = bbox.get("width", 0)
                h = bbox.get("height", 0)
                
                crop_box = (x, y, x + w, y + h)
                crop_img = img.crop(crop_box)
                
                crop_path = self.crops_dir / f"{item['id']}_crop.jpg"
                crop_img.save(crop_path)
                item["screenshot_path"] = str(crop_path)
        except Exception as e:
            print(f"Failed to generate crops: {e}")

    def _generate_json(self, findings):
        for rtype in ["Executive", "Technical", "Court"]:
            report_findings = self._filter_by_type(findings, rtype)
            out_path = self.reports_dir / f"{rtype.lower()}_report.json"
            with open(out_path, "w") as f:
                json.dump({
                    "report_type": rtype,
                    "findings": report_findings
                }, f, indent=2)

    def _filter_by_type(self, findings, rtype):
        if rtype == "Executive":
            return [f for f in findings if f.get("severity") in ["Critical", "High"]]
        return findings

    def _generate_pdf(self, findings, rtype):
        if not SimpleDocTemplate:
            return
            
        report_findings = self._filter_by_type(findings, rtype)
        out_path = self.reports_dir / f"{rtype.lower()}_report.pdf"
        
        doc = SimpleDocTemplate(str(out_path), pagesize=letter)
        styles = getSampleStyleSheet()
        elements = []
        
        # Title
        elements.append(Paragraph(f"{rtype} Forensic Report", styles['Title']))
        elements.append(Spacer(1, 12))
        
        if rtype == "Court":
            elements.append(Paragraph("Chain of Custody & Legal Boilerplate:", styles['Heading3']))
            elements.append(Paragraph("This document certifies that the digital forensic analysis was conducted using AI-FORGE automated algorithms. No manual tampering of the digital evidence occurred during analysis.", styles['Normal']))
            elements.append(Spacer(1, 12))
            
        for item in report_findings:
            elements.append(Paragraph(f"Finding ID: {item.get('id', 'N/A')} - {item.get('label', '')}", styles['Heading2']))
            
            data = [
                ["Page", str(item.get("page", 1))],
                ["Severity", item.get("severity", "")],
                ["Confidence", f"{round(item.get('confidence', 0)*100)}%"],
                ["Bounding Box", f"x:{item['bbox'].get('x')}, y:{item['bbox'].get('y')}, w:{item['bbox'].get('width')}, h:{item['bbox'].get('height')}"],
                ["Reason", item.get("reason", "")],
                ["Recommendation", item.get("recommendation", "")]
            ]
            
            t = Table(data, colWidths=[1.5*inch, 4.5*inch])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (0,-1), colors.lightgrey),
                ('GRID', (0,0), (-1,-1), 1, colors.black),
                ('VALIGN', (0,0), (-1,-1), 'TOP')
            ]))
            elements.append(t)
            elements.append(Spacer(1, 12))
            
            # Screenshot
            crop_path = item.get("screenshot_path")
            if crop_path and os.path.exists(crop_path):
                elements.append(Paragraph("Visual Evidence:", styles['Heading3']))
                # Try to load and scale image nicely
                try:
                    img = RLImage(crop_path, width=3*inch, height=1.5*inch, kind='proportional')
                    elements.append(img)
                except:
                    pass
                elements.append(Spacer(1, 12))
                
        doc.build(elements)

    def _generate_docx(self, findings, rtype):
        if not docx:
            return
            
        report_findings = self._filter_by_type(findings, rtype)
        out_path = self.reports_dir / f"{rtype.lower()}_report.docx"
        
        doc = docx.Document()
        doc.add_heading(f"{rtype} Forensic Report", 0)
        
        if rtype == "Court":
            doc.add_heading('Chain of Custody & Legal Boilerplate', level=2)
            doc.add_paragraph("This document certifies that the digital forensic analysis was conducted using AI-FORGE automated algorithms. No manual tampering of the digital evidence occurred during analysis.")
            
        for item in report_findings:
            doc.add_heading(f"Finding ID: {item.get('id', 'N/A')} - {item.get('label', '')}", level=1)
            
            table = doc.add_table(rows=6, cols=2)
            table.style = 'Table Grid'
            
            cells = table.rows[0].cells
            cells[0].text = "Page"
            cells[1].text = str(item.get("page", 1))
            
            cells = table.rows[1].cells
            cells[0].text = "Severity"
            cells[1].text = item.get("severity", "")
            
            cells = table.rows[2].cells
            cells[0].text = "Confidence"
            cells[1].text = f"{round(item.get('confidence', 0)*100)}%"
            
            cells = table.rows[3].cells
            cells[0].text = "Bounding Box"
            cells[1].text = f"x:{item['bbox'].get('x')}, y:{item['bbox'].get('y')}, w:{item['bbox'].get('width')}, h:{item['bbox'].get('height')}"
            
            cells = table.rows[4].cells
            cells[0].text = "Reason"
            cells[1].text = item.get("reason", "")
            
            cells = table.rows[5].cells
            cells[0].text = "Recommendation"
            cells[1].text = item.get("recommendation", "")
            
            crop_path = item.get("screenshot_path")
            if crop_path and os.path.exists(crop_path):
                doc.add_heading('Visual Evidence', level=2)
                try:
                    doc.add_picture(crop_path, width=Inches(3))
                except:
                    pass
                    
            doc.add_page_break()
            
        doc.save(str(out_path))

if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        engine = CourtReportEngine(original_image_path=sys.argv[1])
        engine.generate_reports()
